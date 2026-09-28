# -*- coding: utf-8 -*-
"""
Genera la documentación de variables a partir de datos_variables.py:
  - docs/01_Tabla_Variables.html
  - tia/Variables_PLC.xlsx   (formato de importación de variables PLC de TIA Portal V15)

Uso:  python3 herramientas/generar.py      (requiere: pip install openpyxl)
"""
import html
import os
import re
import sys

sys.path.insert(0, os.path.dirname(__file__))
import datos_variables as D  # noqa: E402

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def e(texto):
    """Escapa HTML y evita que '%' o 'L' se separen de su número al cortar la línea."""
    t = html.escape(texto).replace(" %", "&nbsp;%")
    return re.sub(r"(\d) L\b", r"\1&nbsp;L", t)


# ---------------------------------------------------------------------------
# Excel importable en TIA Portal
# ---------------------------------------------------------------------------
def generar_excel(ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "PLC Tags"
    cab = ["Name", "Path", "Data Type", "Logical Address", "Comment",
           "Hmi Visible", "Hmi Accessible", "Hmi Writeable", "Typeobject ID", "Version ID"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    filas = [(n, t, a, c) for (a, n, t, c, _x, _f) in D.ENTRADAS_DIGITALES if n]
    filas += [(n, t, a, c) for (a, n, t, c, _x, _f) in D.SALIDAS_DIGITALES]
    filas += [(n, t, a, c) for (a, n, t, c, _x, _f) in D.ANALOGICAS]
    for n, t, a, c in filas:
        ws.append([n, D.TABLA_TIA, t, a, c, "True", "True", "True", "", ""])
    for col, ancho in zip("ABCDEFGHIJ", (26, 16, 10, 16, 52, 12, 14, 14, 14, 11)):
        ws.column_dimensions[col].width = ancho
    wb.create_sheet("User Constants").append(
        ["Name", "Path", "Data Type", "Value", "Comment"])
    wb.save(ruta)
    return len(filas)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------
def addr_chip(a):
    if not a:
        return ""
    area = "in" if a.startswith("%I") else "out" if a.startswith("%Q") else "mem"
    return f'<span class="addr addr-{area}">{e(a)}</span>'


def tabla(cabeceras, filas, clase=""):
    th = "".join(f"<th>{h}</th>" for h in cabeceras)
    return (f'<div class="tw"><table class="vt {clase}"><thead><tr>{th}</tr></thead>'
            f'<tbody>{"".join(filas)}</tbody></table></div>')


def filas_io(lista):
    out = []
    for a, n, t, c, x, f in lista:
        if not n:
            out.append(f'<tr class="reserva"><td>{addr_chip(a)}</td><td colspan="5">'
                       f'Reserva{(" · " + e(x)) if x else ""}</td></tr>')
            continue
        out.append(
            f"<tr><td>{addr_chip(a)}</td><td class='nm'>{e(n)}</td><td class='ty'>{e(t)}</td>"
            f"<td class='cm'>{e(c)}</td><td>{e(x)}</td><td class='fio'>{e(f)}</td></tr>")
    return out


LAMP = {"V": "g", "N": "o", "R": "r", "Verde": "g", "Naranja": "o", "Rojo": "r"}


def lamp(c, blink=False):
    return f'<span class="lamp lamp-{LAMP[c]}{" blink" if blink else ""}" aria-hidden="true"></span>'


def seccion(id_, titulo, intro, cuerpo):
    return (f'<section id="{id_}"><h2>{titulo}</h2>'
            f'{f"<p class=lead>{intro}</p>" if intro else ""}{cuerpo}</section>')


ESQUEMA = """
<div class="tw schem"><svg viewBox="0 0 980 560" role="img" aria-labelledby="sch-t">
<title id="sch-t">Esquema de la línea con la posición de cada sensor y actuador</title>
<defs>
 <pattern id="belt" width="14" height="22" patternUnits="userSpaceOnUse">
  <rect width="14" height="22" class="s-belt"/><line x1="0" y1="0" x2="0" y2="22" class="s-beltln"/>
 </pattern>
</defs>
<!-- bomba y válvula de seguridad -->
<line x1="385" y1="24" x2="385" y2="112" class="s-pipe"/>
<circle cx="385" cy="34" r="16" class="s-dev"/><path d="M377 42 L385 24 L393 42 Z" class="s-devin"/>
<text x="410" y="30" class="s-lbl">Bomba</text><text x="410" y="44" class="s-addr">QW64</text>
<path d="M371 74 L399 90 L399 74 L371 90 Z" class="s-valve"/>
<text x="410" y="80" class="s-lbl">Válv. seguridad</text><text x="410" y="94" class="s-addr">Q0.0</text>
<!-- tanque -->
<rect x="330" y="112" width="110" height="120" rx="4" class="s-tank"/>
<rect x="332" y="160" width="106" height="70" rx="2" class="s-water"/>
<line x1="314" y1="202" x2="330" y2="202" class="s-tick"/><text x="308" y="206" class="s-addr" text-anchor="end">25 %</text>
<line x1="314" y1="124" x2="330" y2="124" class="s-tick"/><text x="308" y="136" class="s-addr" text-anchor="end">90 %</text>
<line x1="314" y1="114" x2="330" y2="114" class="s-tick"/><text x="308" y="112" class="s-addr" text-anchor="end">99 %</text>
<text x="308" y="160" class="s-note" text-anchor="end">Nivel IW64</text>
<text x="385" y="180" class="s-lbl" text-anchor="middle">Tanque</text>
<!-- descarga y caudalímetro -->
<line x1="385" y1="232" x2="385" y2="282" class="s-pipe"/>
<path d="M371 244 L399 260 L399 244 L371 260 Z" class="s-valve"/>
<text x="408" y="250" class="s-lbl">Válv. llenado</text><text x="408" y="264" class="s-addr">QW66</text>
<rect x="377" y="266" width="16" height="12" class="s-dev"/>
<text x="364" y="277" class="s-addr" text-anchor="end">Caudal IW66</text>
<!-- cinta principal -->
<rect x="8" y="300" width="84" height="44" rx="4" class="s-dev"/>
<text x="50" y="326" class="s-lbl" text-anchor="middle">Generador</text>
<text x="50" y="364" class="s-addr" text-anchor="middle">Q0.3</text>
<rect x="96" y="312" width="214" height="22" fill="url(#belt)" class="s-beltbox"/>
<rect x="316" y="312" width="140" height="22" fill="url(#belt)" class="s-beltbox"/>
<rect x="462" y="312" width="140" height="22" fill="url(#belt)" class="s-beltbox"/>
<rect x="608" y="312" width="200" height="22" fill="url(#belt)" class="s-beltbox"/>
<text x="203" y="354" class="s-cv" text-anchor="middle">C1 · Q0.4</text>
<text x="386" y="354" class="s-cv" text-anchor="middle">C2 · Q0.5</text>
<text x="532" y="354" class="s-cv" text-anchor="middle">C3 · Q0.6</text>
<text x="650" y="354" class="s-cv" text-anchor="middle">C4 · Q0.7</text>
<!-- flujo -->
<path d="M110 390 L790 390" class="s-flow"/><path d="M790 390 l-8 -5 v10 z" class="s-flowhd"/>
<text x="110" y="408" class="s-note">sentido de avance · hasta 4 botellas a la vez (FIFO)</text>
<!-- detección -->
<line x1="200" y1="270" x2="200" y2="312" class="s-beam"/>
<circle cx="200" cy="296" r="5" class="s-sen"/><circle cx="200" cy="278" r="5" class="s-sen"/>
<text x="192" y="282" class="s-addr" text-anchor="end">Alto I1.1</text>
<text x="192" y="300" class="s-addr" text-anchor="end">Bajo I1.0</text>
<text x="200" y="258" class="s-lbl" text-anchor="middle">Detección</text>
<!-- sensores de posición -->
<circle cx="440" cy="306" r="5" class="s-sen"/><text x="440" y="296" class="s-addr" text-anchor="middle">I1.2</text>
<circle cx="586" cy="306" r="5" class="s-sen"/><text x="586" y="296" class="s-addr" text-anchor="middle">I1.3</text>
<circle cx="740" cy="306" r="5" class="s-sen"/><text x="758" y="300" class="s-addr">I1.4</text>
<!-- TAP-01 -->
<rect x="506" y="222" width="54" height="36" rx="3" class="s-dev"/>
<text x="533" y="244" class="s-lbl" text-anchor="middle">TAP-01</text>
<line x1="533" y1="258" x2="533" y2="306" class="s-rod"/>
<text x="533" y="210" class="s-addr" text-anchor="middle">Q2.0 · I2.0/I2.1</text>
<!-- empujador y caja malas -->
<rect x="692" y="252" width="44" height="26" rx="3" class="s-dev"/>
<line x1="714" y1="278" x2="714" y2="306" class="s-rod"/>
<text x="714" y="226" class="s-lbl" text-anchor="middle">Empujador</text>
<text x="714" y="242" class="s-addr" text-anchor="middle">Q2.1 · I2.2/I2.3</text>
<rect x="654" y="430" width="120" height="52" rx="3" class="s-box"/>
<text x="716" y="452" class="s-lbl" text-anchor="middle">MALAS</text>
<text x="716" y="468" class="s-addr" text-anchor="middle">20 · I2.4/I3.0</text>
<path d="M714 340 L714 424" class="s-reject"/><path d="M714 424 l-5 -8 h10 z" class="s-rejecthd"/>
<!-- transversal -->
<rect x="814" y="146" width="30" height="334" fill="url(#belt)" class="s-beltbox"/>
<circle cx="822" cy="323" r="5" class="s-sen"/><text x="852" y="318" class="s-addr">I1.5</text>
<text x="852" y="304" class="s-cv">C5</text>
<text x="852" y="258" class="s-addr">Q1.0 ↑</text><text x="852" y="392" class="s-addr">Q1.1 ↓</text>
<rect x="769" y="84" width="120" height="54" rx="3" class="s-box"/>
<text x="829" y="106" class="s-lbl" text-anchor="middle">0,5 L</text>
<text x="829" y="122" class="s-addr" text-anchor="middle">6 · I2.5/I3.1</text>
<rect x="769" y="488" width="120" height="54" rx="3" class="s-box"/>
<text x="829" y="510" class="s-lbl" text-anchor="middle">2 L</text>
<text x="829" y="526" class="s-addr" text-anchor="middle">6 · I2.6/I3.2</text>
</svg></div>
"""


def construir_html():
    # --- hardware
    hw = tabla(["Módulo", "Referencia", "Direcciones", "Qué conecta"],
               [f"<tr><td class='nm'>{e(m)}</td><td class='ty'>{e(r)}</td>"
                f"<td class='mono'>{e(d)}</td><td>{e(q)}</td></tr>" for m, r, d, q in D.HARDWARE])
    hw += ("<p class='note'>Al añadir cada módulo en <b>Configuración de dispositivos</b>, cambia su "
           "dirección inicial a la indicada. Así las direcciones quedan seguidas y el driver de "
           "Factory I/O (S7-PLCSIM) las lee en un único bloque.</p>")

    cab_io = ["Dirección", "Nombre", "Tipo", "Comentario TIA", "Qué hace", "Factory I/O"]
    di = tabla(cab_io, filas_io(D.ENTRADAS_DIGITALES))

    salidas_sin_lamp = [s for s in D.SALIDAS_DIGITALES if not s[1].startswith("L")]
    dq = tabla(cab_io, filas_io(salidas_sin_lamp))

    # lámparas: matriz fase x color
    filas_l = []
    por_fase = {}
    for fase, color, addr, nombre in D.LAMPARAS:
        por_fase.setdefault(fase, {})[color] = (addr, nombre)
    for fase in sorted(por_fase):
        celdas = "".join(
            f"<td>{lamp(c)}{addr_chip(por_fase[fase][c][0])}"
            f"<span class='nm small'>{e(por_fase[fase][c][1])}</span></td>"
            for c in ("Verde", "Naranja", "Rojo"))
        filas_l.append(f"<tr><td class='fase'><b>{fase}</b> {e(D.FASES_TEXTO[fase])}</td>{celdas}</tr>")
    lamps = tabla(["Fase", "Verde", "Naranja", "Rojo"], filas_l, "lamps")

    an = tabla(cab_io, filas_io(D.ANALOGICAS))
    an += ("<p class='note'>Escalado: <span class='mono'>NORM_X</span> (0…27648 → 0.0…1.0) y "
           "<span class='mono'>SCALE_X</span> (0.0…1.0 → 0…100 % o 0…Caudal_Max). "
           "Para la salida, al revés: valor en % → <span class='mono'>NORM_X</span> → "
           "<span class='mono'>SCALE_X</span> 0…27648.</p>")

    mk = tabla(["Dirección", "Nombre", "Tipo", "Comentario", "Para qué"],
               [f"<tr><td>{addr_chip(a)}</td><td class='nm'>{e(n)}</td><td class='ty'>{e(t)}</td>"
                f"<td class='cm'>{e(c)}</td><td>{e(x)}</td></tr>" for a, n, t, c, x in D.MARCAS])
    mk += ("<p class='note'>Estas no se importan. Se activan en <b>CPU → Propiedades → "
           "Marcas de sistema y de ciclo</b>: byte de marcas de ciclo = 0 (MB0) y byte de marcas "
           "de sistema = 1 (MB1). TIA crea las variables solo, con estos nombres.</p>")

    # DB_Linea agrupada
    filas_db, grupo_prev = [], None
    for g, n, t, v, r, x in D.DB_LINEA:
        if g != grupo_prev:
            filas_db.append(f"<tr class='grp'><td colspan='5'>{e(g)} <span>Struct</span></td></tr>")
            grupo_prev = g
        filas_db.append(
            f"<tr><td class='nm'>{e(g)}.{e(n)}</td><td class='ty'>{e(t)}</td>"
            f"<td class='mono'>{e(v)}</td><td>{'<span class=ret>Remanente</span>' if r else ''}</td>"
            f"<td>{e(x)}</td></tr>")
    dbl = tabla(["Variable", "Tipo", "Valor inicial", "Remanencia", "Qué es"], filas_db)

    udt = tabla(["Campo", "Tipo", "Inicial", "Qué es"],
                [f"<tr><td class='nm'>{e(n)}</td><td class='ty'>{e(t)}</td><td class='mono'>{e(v)}</td>"
                 f"<td>{e(x)}</td></tr>" for n, t, v, x in D.UDT_BOTELLA])
    fifo = tabla(["Variable", "Tipo", "Inicial", "Qué es"],
                 [f"<tr><td class='nm'>{e(n)}</td><td class='ty'>{e(t)}</td><td class='mono'>{e(v)}</td>"
                  f"<td>{e(x)}</td></tr>" for n, t, v, x in D.DB_FIFO])
    fifo_txt = (
        "<div class='fifo'>"
        "<ol class='steps'>"
        "<li><b>Emisión.</b> Sale una botella: se guarda en <span class='mono'>Botella[Ptr_Entrada]</span> con "
        "<span class='mono'>Activa = 1</span> y la hora (<span class='mono'>T_Emision</span>). El puntero avanza.</li>"
        "<li><b>Detección.</b> Los sensores de altura escriben <span class='mono'>Es_2L</span> en "
        "<span class='mono'>Botella[Ptr_Deteccion]</span>.</li>"
        "<li><b>Llenado.</b> Lee el tipo para calcular la consigna y escribe <span class='mono'>Mala</span> y "
        "<span class='mono'>Volumen_L</span> en <span class='mono'>Botella[Ptr_Llenado]</span>.</li>"
        "<li><b>Rechazo.</b> Lee <span class='mono'>Mala</span> de <span class='mono'>Botella[Ptr_Rechazo]</span>. "
        "Si es mala, la empuja y la borra (<span class='mono'>Activa = 0</span>).</li>"
        "<li><b>Clasificación.</b> La buena pasa a C5, se lee su tipo y, al caer en su caja, se borra.</li>"
        "</ol>"
        "<p class='note'>Cada estación tiene su propio puntero porque las botellas nunca se adelantan. "
        "Una botella nueva solo se emite si <span class='mono'>N_Botellas &lt; Max_Botellas</span>, ha pasado "
        "<span class='mono'>T_Espera_Emision</span> y la posición <span class='mono'>Botella[Ptr_Entrada]</span> "
        "está libre.</p></div>")

    tmp = tabla(["Instancia", "Tipo", "Tiempo / PV", "Fase", "Para qué"],
                [f"<tr><td class='nm'>{e(n)}</td><td class='ty'>{e(t)}</td><td class='mono'>{e(p)}</td>"
                 f"<td>{e(f)}</td><td>{e(x)}</td></tr>" for n, t, p, f, x in D.TEMPORIZADORES])

    # alarmas
    leyenda = (
        "<div class='legend'>"
        f"<div>{lamp('V')}<b>Verde fijo</b><span>Fase trabajando bien</span></div>"
        f"<div>{lamp('V', True)}<b>Verde parpadeo</b><span>Fase en espera</span></div>"
        f"<div>{lamp('N')}<b>Naranja</b><span>Aviso, la línea sigue</span></div>"
        f"<div>{lamp('R')}<b>Rojo</b><span>Fallo, la línea se para</span></div>"
        "</div>")
    ver = tabla(["Fase", "Verde fijo", "Verde parpadeo"],
                [f"<tr><td class='fase'><b>{f}</b> {e(D.FASES_TEXTO[f])}</td><td>{e(a)}</td><td>{e(b)}</td></tr>"
                 for f, a, b in D.VERDES])
    al = tabla(["ID", "Fase", "Color", "Alarma", "Cuándo salta", "Qué hace la línea", "Cómo se quita"],
               [f"<tr class='al-{LAMP[c]}'><td class='mono'><b>{i}</b></td><td>{f}</td><td>{lamp(c)}</td>"
                f"<td class='nm2'>{e(t)}</td><td>{e(cond)}</td><td>{e(q)}</td><td>{e(r)}</td></tr>"
                for i, f, c, t, cond, q, r in D.ALARMAS])

    bloques = tabla(["Bloque", "Nombre", "Qué hace"],
                    [f"<tr><td class='mono'><b>{e(b)}</b></td><td class='nm'>{e(n)}</td><td>{e(x)}</td></tr>"
                     for b, n, x in D.BLOQUES])

    pend = "<ol class='pend'>" + "".join(
        f"<li><b>{e(t)}</b><p>{e(x)}</p></li>" for t, x in D.PENDIENTES) + "</ol>"

    n_di = sum(1 for x in D.ENTRADAS_DIGITALES if x[1])
    n_dq = len(D.SALIDAS_DIGITALES)

    nav = [("hw", "Hardware"), ("esquema", "Esquema"), ("di", "Entradas"), ("dq", "Salidas"),
           ("lamparas", "Lámparas"), ("an", "Analógicas"), ("marcas", "Marcas"), ("db", "DB_Linea"),
           ("fifo", "FIFO"), ("tiempos", "Tiempos"), ("alarmas", "Alarmas"), ("bloques", "Bloques"),
           ("pendiente", "Pendiente")]
    nav_html = "".join(f"<a href='#{i}'>{t}</a>" for i, t in nav)

    cuerpo = "".join([
        seccion("hw", "Hardware y direccionamiento",
                "Proyecto en TIA Portal V15 para S7-1200, simulado con S7-PLCSIM y Factory I/O.", hw),
        seccion("esquema", "Esquema de la línea",
                "Dónde está cada sensor y cada actuador. Las direcciones coinciden con las tablas.", ESQUEMA),
        seccion("di", "Entradas digitales",
                f"{n_di} entradas en uso. Los pulsadores de PARO y SETA son NC: en reposo están a 1.", di),
        seccion("dq", "Salidas digitales",
                f"{n_dq} salidas en uso, contando las 21 lámparas de fase de la sección siguiente.", dq),
        seccion("lamparas", "Lámparas por fase",
                "Tres lámparas por fase, de Q3.3 a Q5.7.", lamps),
        seccion("an", "Entradas y salidas analógicas", "Rango Siemens estándar: 0…27648 = 0…10 V.", an),
        seccion("marcas", "Marcas de sistema y de ciclo", "", mk),
        seccion("db", "DB1 · DB_Linea",
                "Bloque de datos global con acceso optimizado. Todo lo que no es una E/S física está aquí, "
                "agrupado por Struct. Los parámetros son remanentes: no se pierden al apagar la CPU.", dbl),
        seccion("fifo", "UDT_Botella y DB2 · DB_FIFO",
                "Cada botella lleva sus datos desde que sale del generador hasta que cae en una caja.",
                "<h3>UDT_Botella</h3>" + udt + "<h3>DB_FIFO</h3>" + fifo + "<h3>Cómo se recorre</h3>" + fifo_txt),
        seccion("tiempos", "DB3 · DB_Tiempos",
                "Temporizadores y contadores IEC en un DB global, sin mezclarlos dentro de las FC.", tmp),
        seccion("alarmas", "Alarmas por fase", "Cada fase tiene lámpara verde, naranja y roja.",
                leyenda + "<h3>Significado del verde</h3>" + ver + "<h3>Lista de alarmas</h3>" + al),
        seccion("bloques", "Estructura de bloques", "Orden de llamada dentro del OB1: de FC1 a FC10.", bloques),
        seccion("pendiente", "Pendiente de confirmar",
                "Con esto contestado paso a programar bloque por bloque.", pend),
    ])

    with open(os.path.join(os.path.dirname(__file__), "plantilla.html"), encoding="utf-8") as fh:
        plantilla = fh.read()
    return (plantilla
            .replace("{{NAV}}", nav_html)
            .replace("{{CUERPO}}", cuerpo)
            .replace("{{REV}}", e(D.REVISION))
            .replace("{{FECHA}}", e(D.FECHA))
            .replace("{{N_DI}}", str(n_di))
            .replace("{{N_DQ}}", str(n_dq))
            .replace("{{N_AL}}", str(len(D.ALARMAS))))


if __name__ == "__main__":
    pagina = construir_html()
    ruta_html = os.path.join(RAIZ, "docs", "01_Tabla_Variables.html")
    with open(ruta_html, "w", encoding="utf-8") as fh:
        fh.write('<!doctype html>\n<html lang="es">\n<head>\n<meta charset="utf-8">\n'
                 '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
                 '</head>\n<body>\n' + pagina + '\n</body>\n</html>\n')
    print("HTML  ->", ruta_html)
    # Versión sin cabecera <html>, para publicarla como página web (Artifact)
    if len(sys.argv) > 1:
        with open(sys.argv[1], "w", encoding="utf-8") as fh:
            fh.write(pagina)
        print("Web   ->", sys.argv[1])
    ruta_xlsx = os.path.join(RAIZ, "tia", "Variables_PLC.xlsx")
    n = generar_excel(ruta_xlsx)
    print(f"Excel -> {ruta_xlsx} ({n} variables)")
